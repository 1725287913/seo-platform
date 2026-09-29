# knowledge.py
# 知识库（轻量本地 RAG）：把资料文本存本地，出稿时按关键词检索相关段落注入 prompt。
# 不依赖向量库/嵌入模型 —— 用「中文 2-gram + 命中密度」打分，零成本、可离线。
import json, os, re, time, uuid

DATA = os.path.join(os.path.dirname(__file__), "data", "knowledge.json")


def _read():
    if not os.path.exists(DATA):
        json.dump([], open(DATA, "w", encoding="utf-8"))
        return []
    return json.load(open(DATA, encoding="utf-8"))


def _write(arr):
    json.dump(arr, open(DATA, "w", encoding="utf-8"), ensure_ascii=False, indent=2)


def list_docs():
    """只返回元信息（不返回全文），列表页不卡。"""
    return [{"id": d["id"], "title": d["title"], "tags": d.get("tags", ""),
             "chars": len(d.get("content", "")), "created_at": d.get("created_at", "")}
            for d in _read()]


def get_doc(did):
    return next((d for d in _read() if d.get("id") == did), None)


def add_doc(d: dict):
    arr = _read()
    title = (d.get("title") or "").strip()
    content = (d.get("content") or "").strip()
    if not title or not content:
        return {"error": "empty", "message": "标题和内容都不能为空"}
    item = {"id": str(uuid.uuid4())[:8], "title": title, "content": content,
            "tags": d.get("tags", ""), "created_at": time.strftime("%Y-%m-%d %H:%M")}
    arr.insert(0, item)
    _write(arr)
    return item


def del_doc(did):
    _write([d for d in _read() if d.get("id") != did])
    return {"ok": True}


# ---------- 文件导入（PDF / Word / Excel / 纯文本）----------

MAX_DOC_CHARS = 15000   # 单篇超长就切份，检索更准、前端打开不卡


def parse_file(filename: str, data: bytes) -> str:
    """按扩展名把上传的文件解析成纯文本。"""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        import io
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".docx":
        import io
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for tb in doc.tables:                      # 表格里的资料也要抽出来
            for row in tb.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    if ext == ".xlsx":
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"【工作表：{ws.title}】")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    # txt / md / csv / json / html 等纯文本：utf-8 优先，中文老文件多为 gbk
    for enc in ("utf-8", "gbk", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _clean(text: str) -> str:
    """清理 PDF 常见的排版噪声：压缩空格与连续空行。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split(text: str, size: int):
    """按段落累积切分，尽量不切断段落。"""
    if len(text) <= size:
        return [text]
    chunks, buf = [], ""
    for para in text.split("\n"):
        if buf and len(buf) + len(para) + 1 > size:
            chunks.append(buf.strip())
            buf = ""
        buf += para + "\n"
    if buf.strip():
        chunks.append(buf.strip())
    return chunks


def import_file(filename: str, data: bytes, tags: str = ""):
    """导入文件 → 解析成文本 → 超长自动切份入库。"""
    try:
        text = _clean(parse_file(filename, data))
    except Exception as e:
        return {"error": "parse_failed", "message": f"解析失败：{e}"}
    if len(text) < 10:
        return {"error": "empty",
                "message": "没解析出文字。若是扫描版 PDF（图片型），需要先 OCR 成文字"}
    base = os.path.splitext(os.path.basename(filename))[0]
    chunks = _split(text, MAX_DOC_CHARS)
    created = []
    for i, ck in enumerate(chunks, 1):
        title = base if len(chunks) == 1 else f"{base}（{i}/{len(chunks)}）"
        created.append(add_doc({"title": title, "content": ck, "tags": tags}))
    return {"ok": True, "count": len(created), "chars": len(text),
            "titles": [c["title"] for c in created]}


# ---------- 检索部分 ----------

def _tokens(q: str):
    """把查询切成检索单元：中文按 2-gram，英文/数字按小写词。"""
    q = (q or "").lower()
    grams = set()
    for block in re.findall(r"[\u4e00-\u9fa5]+", q):
        # 中文长串切成 2 字滑动窗口（"南京油价" → 南京/京油/油价）
        for i in range(len(block) - 1):
            grams.add(block[i:i + 2])
        if len(block) == 1:
            grams.add(block)
    for w in re.findall(r"[a-z0-9]{2,}", q):
        grams.add(w)
    return grams


def _paragraphs(text: str):
    """按空行/换行切段落，太短的丢弃。"""
    parts = re.split(r"\n\s*\n|\n", text)
    return [p.strip() for p in parts if len(p.strip()) >= 10]


def search(query: str, top: int = 3):
    """在知识库里检索与 query 最相关的段落。
    返回 [{doc_title, text, score}]，score 为命中密度（越高越相关）。
    """
    grams = _tokens(query)
    if not grams:
        return []
    hits = []
    for doc in _read():
        for para in _paragraphs(doc.get("content", "")):
            low = para.lower()
            hit = sum(1 for g in grams if g in low)
            if not hit:
                continue
            # 命中数为主，段落越长越稀释（避免整段长文永远压过精准短句）
            score = hit / (len(para) ** 0.5)
            hits.append({"doc_title": doc["title"], "text": para,
                         "score": round(score, 4), "hit": hit})
    hits.sort(key=lambda x: (-x["hit"], -x["score"]))
    return hits[:top]


def build_context(query: str, top: int = 3, max_chars: int = 1500):
    """把检索结果拼成一段可直接塞进 prompt 的参考资料文本。"""
    hits = search(query, top)
    if not hits:
        return ""
    buf = ["【以下为知识库中检索到的真实资料，请优先依据它作答】"]
    total = 0
    for h in hits:
        piece = f"- （来自《{h['doc_title']}》）{h['text']}"
        if total + len(piece) > max_chars:
            break
        buf.append(piece)
        total += len(piece)
    return "\n".join(buf)
